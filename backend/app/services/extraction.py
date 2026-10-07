import io
import logging
from typing import List, Dict, Any, Optional
import pypdf
import docx

logger = logging.getLogger(__name__)


class ExtractedPage:
    def __init__(self, page_number: Optional[int], content: str):
        self.page_number = page_number
        self.content = content.strip()

    def to_dict(self) -> Dict[str, Any]:
        return {
            "page_number": self.page_number,
            "content": self.content,
        }


def extract_text_from_pdf(file_bytes: bytes) -> List[ExtractedPage]:
    pages = []
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(ExtractedPage(page_number=idx, content=text))
    return pages


def extract_text_from_docx(file_bytes: bytes) -> List[ExtractedPage]:
    doc = docx.Document(io.BytesIO(file_bytes))
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    full_text = "\n\n".join(text_parts)
    return [ExtractedPage(page_number=1, content=full_text)] if full_text.strip() else []


def extract_text_from_txt(file_bytes: bytes) -> List[ExtractedPage]:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")

    return [ExtractedPage(page_number=1, content=text)] if text.strip() else []


def extract_text_from_markdown(file_bytes: bytes) -> List[ExtractedPage]:
    return extract_text_from_txt(file_bytes)


def extract_document_content(file_bytes: bytes, file_type: str) -> List[ExtractedPage]:
    """
    Extracts text sections and pages based on file extension/type.
    """
    ft = file_type.lower().strip(".")
    if ft == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ft in ("docx", "doc"):
        return extract_text_from_docx(file_bytes)
    elif ft in ("txt", "text"):
        return extract_text_from_txt(file_bytes)
    elif ft in ("md", "markdown"):
        return extract_text_from_markdown(file_bytes)
    else:
        # Fallback to UTF-8 plain text attempt
        return extract_text_from_txt(file_bytes)
