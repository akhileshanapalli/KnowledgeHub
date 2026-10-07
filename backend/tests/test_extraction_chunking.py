import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
import docx
from app.main import app
from app.db.session import get_db
from app.core.security import create_access_token
from app.core.rate_limit import reset_rate_limit_store
from app.models.user import User
from app.models.workspace import Workspace, WorkspaceMembership
from app.models.document import Document, DocumentChunk
from app.services.extraction import (
    extract_document_content,
    extract_text_from_txt,
    extract_text_from_markdown,
    extract_text_from_docx,
    ExtractedPage,
)
from app.services.chunking import split_text_recursively, chunk_extracted_pages


@pytest.fixture
def client(db_session: Session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    reset_rate_limit_store()
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def test_txt_and_md_extraction():
    sample_txt = b"Section 1: Data Protection\nAll customer data must be encrypted at rest and in transit."
    pages_txt = extract_document_content(sample_txt, "txt")
    assert len(pages_txt) == 1
    assert "Data Protection" in pages_txt[0].content
    assert pages_txt[0].page_number == 1

    sample_md = b"# Welcome to KnowledgeHub\n\nThis is an **enterprise** RAG platform."
    pages_md = extract_document_content(sample_md, "md")
    assert len(pages_md) == 1
    assert "enterprise" in pages_md[0].content


def test_docx_extraction():
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("KnowledgeHub uses FastAPI, pgvector, and Next.js.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Technology"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    pages = extract_text_from_docx(file_bytes)
    assert len(pages) == 1
    assert "Architecture Overview" in pages[0].content
    assert "pgvector" in pages[0].content


def test_recursive_smart_chunker():
    text = (
        "Paragraph 1 contains general information about the workspace. "
        "It defines user roles and access policies.\n\n"
        "Paragraph 2 discusses encryption standards. AES-256 is strictly enforced for all persistent disks.\n\n"
        "Paragraph 3 covers backup schedules. Full snapshots are taken every 24 hours at midnight UTC."
    )
    pages = [
        ExtractedPage(page_number=1, content=text),
        ExtractedPage(page_number=2, content="Page 2 summary: Incident response SLA is 15 minutes."),
    ]

    # Chunk with small size to test splitting & overlap
    chunks = chunk_extracted_pages(pages, chunk_size=120, chunk_overlap=30)
    assert len(chunks) >= 3

    # Check page numbers are preserved
    page_1_chunks = [c for c in chunks if c.page_number == 1]
    page_2_chunks = [c for c in chunks if c.page_number == 2]
    assert len(page_1_chunks) >= 2
    assert len(page_2_chunks) >= 1
    assert "Incident response" in page_2_chunks[0].content


def test_document_upload_and_ingestion_api(client: TestClient, db_session: Session):
    # Setup User and Workspace
    user = User(email="analyst@corp.com", hashed_password="pw", full_name="Analyst")
    ws = Workspace(name="Corp Research", slug="corp-research")
    db_session.add_all([user, ws])
    db_session.flush()

    membership = WorkspaceMembership(user_id=user.id, workspace_id=ws.id, role="owner")
    db_session.add(membership)
    db_session.commit()

    token = create_access_token(user.id)

    # Upload document
    file_content = (
        b"Company Travel Guidelines:\n"
        b"1. Domestic travel flights must be booked in Economy class.\n"
        b"2. Meals per diem allowance is $75 per day.\n"
        b"3. Hotel accommodations up to $200 per night are reimbursable."
    )
    files = {"file": ("travel_policy.txt", file_content, "text/plain")}
    data = {"source_url": "https://intranet.corp.com/policies/travel"}

    res = client.post(
        "/api/v1/documents/upload",
        headers={"Authorization": f"Bearer {token}", "X-Workspace-ID": ws.id},
        files=files,
        data=data,
    )
    assert res.status_code == 202, res.text
    doc_data = res.json()
    doc_id = doc_data["id"]
    assert doc_data["filename"] == "travel_policy.txt"
    assert doc_data["status"] in ("uploaded", "completed")

    # Verify status endpoint
    status_res = client.get(
        f"/api/v1/documents/{doc_id}/status",
        headers={"Authorization": f"Bearer {token}", "X-Workspace-ID": ws.id},
    )
    assert status_res.status_code == 200
    assert status_res.json()["chunk_count"] >= 1

    # Verify chunks endpoint
    chunks_res = client.get(
        f"/api/v1/documents/{doc_id}/chunks",
        headers={"Authorization": f"Bearer {token}", "X-Workspace-ID": ws.id},
    )
    assert chunks_res.status_code == 200
    chunks = chunks_res.json()
    assert len(chunks) >= 1
    assert "Economy class" in chunks[0]["content"]

    # Verify document deletion
    del_res = client.delete(
        f"/api/v1/documents/{doc_id}",
        headers={"Authorization": f"Bearer {token}", "X-Workspace-ID": ws.id},
    )
    assert del_res.status_code == 204

    # Verify document is gone
    check_res = client.get(
        f"/api/v1/documents/{doc_id}",
        headers={"Authorization": f"Bearer {token}", "X-Workspace-ID": ws.id},
    )
    assert check_res.status_code == 404
